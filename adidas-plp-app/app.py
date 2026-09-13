from io import BytesIO
from flask import Flask, jsonify, render_template, request, send_file, send_from_directory
from werkzeug.exceptions import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from scraper import fetch_page, normalize_url, InputError, AdidasError, MAX_PRODUCTS

app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 3 * 1024 * 1024

COLUMNS = [('id', 'ID', 15), ('name', 'Name', 48), ('subtitle', 'Description', 26),
           ('price', 'Current price USD', 21), ('original_price', 'Original price USD', 22),
           ('discount', 'Discount %', 17), ('currency', 'Currency', 12),
           ('category', 'Category', 18), ('sport', 'Sport', 18), ('division', 'Division', 18),
           ('sizes', 'Listed sizes', 45), ('color_variants', 'Listed variants', 22),
           ('rating', 'Rating', 15), ('reviews', 'Reviews', 15), ('orderable', 'Orderable', 19), ('url', 'Link', 65)]


@app.get('/')
def home():
    return render_template('index.html')


@app.get('/adidas-logo.png')
@app.get('/app.js')
@app.get('/style.css')
def assets():
    return send_from_directory('public', request.path.lstrip('/'))


@app.post('/api/page')
def page():
    body = request.get_json()
    if not isinstance(body, dict):
        raise InputError('Invalid request.')
    return jsonify(fetch_page(body.get('url'), body.get('start', 0)))


def text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    # Datos externos se guardan como texto, nunca como fórmulas de Excel.
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


@app.post('/api/export')
def export():
    data = request.get_json()
    if not isinstance(data, dict):
        raise InputError('Invalid request.')
    url = normalize_url(data.get('url'))
    rows = data.get('products')
    total = data.get('total')
    if not isinstance(rows, list) or type(total) is not int or not 0 <= total <= MAX_PRODUCTS or len(rows) != total:
        raise InputError('Excel export requires the complete list.')
    if any(not isinstance(r, dict) or not isinstance(r.get('id'), str) or not r['id'] for r in rows):
        raise InputError('Invalid product.')
    if len({r['id'] for r in rows}) != total:
        raise InputError('Duplicate products found.')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    for row, label, value in [(1, 'PLP', url), (2, 'Checked at (UTC)', data.get('fetched_at', '')), (3, 'Unique products', total)]:
        text_cell(ws, row, 1, label)
        text_cell(ws, row, 2, str(value) if row != 3 else value)
    for col, (_, name, width) in enumerate(COLUMNS, 1):
        c = ws.cell(5, col, name)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='132A35')
        ws.column_dimensions[c.column_letter].width = width
    for rn, item in enumerate(rows, 6):
        for cn, (key, _, _) in enumerate(COLUMNS, 1):
            value = item.get(key)
            if value is not None and not isinstance(value, (str, int, float)):
                raise InputError('Invalid product data.')
            if isinstance(value, str) and len(value) > 4000:
                raise InputError('Value is too long.')
            c = text_cell(ws, rn, cn, value)
            c.alignment = Alignment(vertical='top')
            if key in ('price', 'original_price'):
                c.number_format = '"$"#,##0.00'
            elif key in ('discount', 'rating'):
                c.number_format = '0.00'
    ws.freeze_panes = 'C6'
    ws.sheet_view.showGridLines = False
    if rows:
        table = Table(displayName='PLPProducts', ref=f'A5:P{5+len(rows)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    ws.row_dimensions[5].height = 25
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='adidas-products.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.errorhandler(InputError)
def input_error(exc):
    return jsonify(error=str(exc)), 400


@app.errorhandler(AdidasError)
def upstream_error(exc):
    return jsonify(error=str(exc)), 502


@app.errorhandler(HTTPException)
def http_error(exc):
    return jsonify(error='Invalid request or request too large.' if exc.code in (400, 413, 415) else exc.description), exc.code


@app.errorhandler(Exception)
def unexpected(exc):
    app.logger.exception('Application error')
    return jsonify(error='Unable to complete the operation. Please try again.'), 500


@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; object-src 'none'; base-uri 'self'; frame-ancestors 'none'"
    if request.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store'
    return response


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5055)
