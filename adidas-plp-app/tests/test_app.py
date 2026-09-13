import unittest
from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
from app import app
from scraper import normalize_url, InputError, product, fetch_page, AdidasError

URL = 'https://www.adidas.com/us/men-running-shoes'

class Tests(unittest.TestCase):
    def setUp(self):
        self.c = app.test_client()

    def test_url_restrictions(self):
        self.assertEqual(normalize_url(URL+'?start=48,'), URL)
        for url in ['https://evil.com/us/shoes', 'https://www.adidas.com.co/us/shoes', URL+'?price=10', 'http://www.adidas.com/us/shoes', 'https://www.adidas.com/us/p/AA.html']:
            with self.assertRaises(InputError): normalize_url(url)

    def test_prices_sizes(self):
        p = product(dict(productId='ABC123',price=100,salePrice=80,availableSizes=['hidden','7','8'],colorVariations=['A','A','B']))
        self.assertEqual((p['price'],p['discount'],p['sizes'],p['color_variants']),(80,20,'7, 8',2))

    def test_bad_json_and_url(self):
        self.assertEqual(self.c.post('/api/page',json=[]).status_code,400)
        self.assertEqual(self.c.post('/api/page',json={'url':'http://localhost'}).status_code,400)

    def test_upstream_block(self):
        with patch('scraper.requests.get') as get:
            get.return_value.status_code=403
            r=self.c.post('/api/page',json={'url':URL})
            self.assertEqual(r.status_code,502)
            self.assertIn('403',r.json['error'])

    def test_repeated_products_rejected(self):
        with patch('scraper.requests.get') as get:
            get.return_value.status_code=200
            get.return_value.json.return_value={'raw':{'itemList':{'count':2,'viewSize':48,'startIndex':0,'items':[{'productId':'X'},{'productId':'X'}]}}}
            with self.assertRaises(AdidasError): fetch_page(URL)

    def test_excel_and_formula_text(self):
        row=product({'productId':'A1','displayName':'=1+1','price':100,'salePrice':80})
        r=self.c.post('/api/export',json={'url':URL,'total':1,'products':[row]})
        self.assertEqual(r.status_code,200)
        ws=load_workbook(BytesIO(r.data)).active
        self.assertEqual(ws['B6'].data_type,'s')
        self.assertEqual(ws['B6'].value,'=1+1')
        self.assertEqual(ws['D6'].value,80)
        self.assertEqual(ws.freeze_panes,'C6')
        self.assertIn('PLPProducts',ws.tables)

    def test_partial_export_rejected(self):
        self.assertEqual(self.c.post('/api/export',json={'url':URL,'total':2,'products':[]}).status_code,400)

    def test_empty_excel(self):
        r=self.c.post('/api/export',json={'url':URL,'total':0,'products':[]})
        self.assertEqual(r.status_code,200)
        self.assertEqual(load_workbook(BytesIO(r.data)).active.max_row,5)

    def test_home_assets(self):
        for path in ['/','/app.js','/style.css']:
            self.assertEqual(self.c.get(path).status_code,200)

if __name__ == '__main__': unittest.main()
