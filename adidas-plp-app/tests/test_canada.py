import unittest
from unittest.mock import patch
from io import BytesIO
from openpyxl import load_workbook
from app import app
from scraper import fetch_page, normalize_url, product, InputError
CA='https://www.adidas.ca/en/men-running-shoes'
class CanadaTests(unittest.TestCase):
 def test_normalize(self):
  self.assertEqual(normalize_url(CA+'?start=96'),CA)
  with self.assertRaises(InputError): normalize_url('https://www.adidas.ca/us/shoes')
 def test_market_mismatch(self):
  with self.assertRaises(InputError): fetch_page(CA,market='US')
 def test_api_mapping_and_export(self):
  with patch('scraper.requests.get') as get:
   get.return_value.status_code=200
   get.return_value.json.return_value={'pageProps':{'title':'Men Running Shoes','sitePath':'en','taxonomy':'men-running-shoes','context':{'locale':'en_CA'},'info':{'count':1,'viewSize':48,'startIndex':0},'products':[{'id':'AA1234','priceData':{'prices':[{'type':'original','value':120}]},'url':'https://www.adidas.ca/en/shoe/AA1234.html'}]}}
   data=fetch_page(CA,market='CA')
   self.assertIn('/plp-app/_next/data/',get.call_args.args[0])
   self.assertTrue(get.call_args.args[0].endswith('/en/men-running-shoes.json'))
   self.assertEqual(get.call_args.kwargs['params'],{'start':0,'path':'en','taxonomy':'men-running-shoes'})
   self.assertEqual(data['products'][0]['currency'],'CAD')
   self.assertTrue(data['products'][0]['url'].startswith('https://www.adidas.ca/en/'))
   with app.test_client() as c:
    response=c.post('/api/export',json=data)
    self.assertEqual(response.status_code,200)
    ws=load_workbook(BytesIO(response.data)).active
    self.assertEqual(ws['D5'].value,'Current price CAD')
    self.assertEqual(ws['G6'].value,'CAD')
    data['products'][0]['currency']='USD'
    self.assertEqual(c.post('/api/export',json=data).status_code,400)

 def test_canadian_soldout_and_missing_rating(self):
  from scraper import canada_product
  p=canada_product({'id':'X1','category':'NOT_IMPLEMENTED','ratingCount':-99,'priceData':{'isSoldOut':True,'prices':[{'type':'original','value':300},{'type':'sale','value':150}]}},[])
  self.assertEqual(p['price'],150)
  self.assertEqual(p['discount'],50)
  self.assertEqual(p['orderable'],'No')
  self.assertIsNone(p['reviews'])
  self.assertEqual(p['division'],'')
 def test_stale_build(self):
  from scraper import AdidasError
  with patch('scraper.requests.get') as get:
   get.return_value.status_code=404
   with self.assertRaisesRegex(AdidasError,'ADIDAS_CA_BUILD_ID'): fetch_page(CA)

 def test_filter_groups(self):
  from scraper import canada_filters
  self.assertIsNone(canada_filters({}))
  self.assertEqual(canada_filters({'filters':[]}),[])
  result=canada_filters({'filters':[{'id':'colour','title':'Colour','values':[{'name':'Black','value':'black','count':12},{'name':'Red','value':'red','count':0}]},{'id':'price_slider','visualization':'slider','values':[{'name':'50','value':'50'},{'name':'400','value':'400'}]}]})
  self.assertEqual(len(result),2)
  self.assertEqual(result[0]['option_count'],2)
  self.assertEqual(result[0]['options'][1]['product_count'],0)
  self.assertEqual(result[1]['kind'],'range')
